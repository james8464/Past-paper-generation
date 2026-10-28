# Build the maintained Prix Occitanie 2026 application evidence documents.

from __future__ import annotations

import argparse
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "occitanie-2026"
LIGHT_GREY = "D9D9D9"
HEADER_GREY = "F2F2F7"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, value: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_borders(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), LIGHT_GREY)


def set_repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    properties.append(cant_split)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_document(
    doc: Document, *, compact: bool = False, editorial: bool = False
) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0 if editorial else 1.7 if compact else 1.9)
    section.bottom_margin = Cm(1.7 if editorial else 1.55 if compact else 1.7)
    section.left_margin = Cm(
        2.1 if editorial and compact else 2.35 if editorial else 1.9 if compact else 2.1
    )
    section.right_margin = Cm(
        2.1 if editorial and compact else 2.35 if editorial else 1.9 if compact else 2.1
    )
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(
        10.5
        if editorial and compact
        else 11
        if editorial
        else 10.2
        if compact
        else 10.8
    )
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(
        5 if editorial and compact else 8 if editorial else 5 if compact else 6
    )
    normal.paragraph_format.line_spacing = (
        1.12 if editorial and compact else 1.22 if editorial else 1.12
    )
    heading_sizes = (
        (("Title", 30), ("Heading 1", 16), ("Heading 2", 12.5))
        if editorial
        else (("Title", 22), ("Heading 1", 13.5), ("Heading 2", 11.5))
    )
    for name, size in heading_sizes:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Title"
        style.paragraph_format.space_before = Pt(
            12 if editorial and compact else 18 if editorial else 12
        )
        style.paragraph_format.space_after = Pt(
            6 if editorial and compact else 9 if editorial else 6
        )
        style.paragraph_format.keep_with_next = True
    title_properties = doc.styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Keldin  ·  Prix Occitanie 2026  ·  ")
    add_page_field(footer)
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(85, 85, 85)


# Use an explicit portable face instead of Word's serif theme fallback.
def finalize_fonts(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.keep_together = True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            if paragraph.style.name == "Title":
                run.font.size = doc.styles["Title"].font.size
                run.font.color.rgb = RGBColor(0, 0, 0)
            elif paragraph.style.name == "Heading 1":
                run.font.size = doc.styles["Heading 1"].font.size
                run.font.bold = True
            elif paragraph.style.name == "Heading 2":
                run.font.size = doc.styles["Heading 2"].font.size
                run.font.bold = True


def title_block(doc: Document, title: str, subtitle: str) -> None:
    title_paragraph = doc.add_paragraph(style="Title")
    title_paragraph.add_run(title)
    subtitle_paragraph = doc.add_paragraph(subtitle)
    subtitle_paragraph.paragraph_format.space_after = Pt(14)
    subtitle_paragraph.runs[0].font.size = Pt(11)
    subtitle_paragraph.runs[0].font.color.rgb = RGBColor(70, 70, 70)


def add_label_paragraph(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(label + " ")
    run.bold = True
    paragraph.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_header(header)
    prevent_row_split(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, HEADER_GREY)
        set_cell_margins(cell, 120)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_margins(cell, 120)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F5F7F9")
            cell.paragraphs[0].add_run(value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    return table


def add_source(doc: Document, number: int, label: str, url: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run(f"[{number}] ").font.size = Pt(9)
    relationship_id = paragraph.part.relate_to(
        url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1D1D1F")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    properties.append(size)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    domain = paragraph.add_run(f"  ·  {urlparse(url).netloc}")
    domain.font.size = Pt(9)
    domain.font.color.rgb = RGBColor(90, 90, 90)


def build_application() -> Path:
    doc = Document()
    configure_document(doc, compact=True)
    title_block(
        doc,
        "Candidature au Prix Occitanie 2026",
        "Keldin pilote NSI Occitanie",
    )
    add_label_paragraph(
        doc, "Candidat", "James Durup · CentraleSupélec · promotion : À compléter"
    )
    add_label_paragraph(
        doc, "Contact", "james.durup@student-cs.fr · téléphone : À compléter"
    )
    add_label_paragraph(doc, "Adresse", "À compléter")
    doc.add_heading("Présentation du candidat", level=1)
    doc.add_paragraph(
        "Je suis étudiant à CentraleSupélec et je développe Keldin, une application macOS de sujets d’entraînement. Ses parcours britanniques utilisent des modèles d’intelligence artificielle. Le parcours français V35 construit cinq familles de sujets de NSI sans appeler de modèle. Je m’intéresse aux logiciels vérifiables et à l’usage responsable de l’IA dans l’éducation. Je présente un prototype dont les questions, réponses et PDF ont fait l’objet d’une évaluation interne comparée aux annales officielles."
    )
    doc.add_heading("Titre du projet", level=1)
    doc.add_paragraph("Keldin pilote NSI Occitanie")
    doc.add_heading("Description synthétique du projet", level=1)
    doc.add_paragraph(
        "Le projet aide à préparer un sujet de Terminale en spécialité Numérique et sciences informatiques, accompagné d’un corrigé proposé et d’un barème indicatif. L’utilisateur choisit l’impression standard ou agrandie, crée un brouillon puis relit les deux PDF. Le catalogue est limité à cinq familles de sujets ; leur répétition est refusée, même après un changement de dossier. Les élèves consultent les PDF sur l’ENT ou sur papier, sans compte ni accès à une IA. L’application ne collecte aucune copie d’élève."
    )
    doc.add_paragraph(
        "La version française reprend la structure annoncée pour 2027 : trois exercices indépendants, 3 h 30, 18 points techniques et 2 points de langue. Chaque question est rattachée à une capacité du programme officiel. Le logiciel vérifie les totaux, les durées estimées, certaines réponses calculables, les requêtes SQL, les traces d’algorithmes, les graphes et la cohérence entre sujet et corrigé. Les PDF sont clairement non officiels. Ils précisent que la difficulté et la durée ne sont pas étalonnées et qu’aucune validation par un enseignant n’est revendiquée."
    )

    doc.add_heading("Originalité du projet", level=1)
    doc.add_paragraph(
        "Le parcours NSI associe des questions, réponses et critères de crédit à des données définies par l’application. Les vérifications portent notamment sur les algorithmes, le SQL et les totaux exacts. Les versions et les empreintes des fichiers restent attachées au sujet. La décision d’un enseignant ne vaut que pour les PDF effectivement relus ; toute modification impose une nouvelle relecture. Ces contrôles ne prouvent pas l’équivalence pédagogique avec le baccalauréat."
    )
    doc.add_paragraph(
        "L’ancrage régional sera construit et mesuré, non simplement nommé. Je souhaite organiser un pilote avec des enseignants des académies de Toulouse et de Montpellier, puis créer un module optionnel fondé sur des données ouvertes d’Occitanie. Les lycées, les transports, la consommation d’énergie et, si les droits le permettent, les ressources en eau pourront nourrir des exercices d’algorithmique et de bases de données sur des questions territoriales et climatiques réelles. Aucun partenariat régional n’est encore conclu."
    )
    doc.add_heading("Bénéficiaires", level=1)
    doc.add_paragraph(
        "Les premiers bénéficiaires sont les enseignants de NSI qui manquent de sujets nouveaux pour entraîner leurs classes. Les élèves reçoivent un PDF ordinaire, lisible sur loRdi, sur l’ENT ou sur papier. La création reste sur le Mac de l’enseignant. Ce choix respecte le parc régional existant : loRdi sert à consulter les ressources, sans prétendre exécuter l’application macOS ou un modèle volumineux."
    )
    doc.add_heading("État de développement", level=1)
    add_bullets(
        doc,
        [
            "Prototype technique : parcours français distinct, génération locale, références filtrées, points décimaux exacts, contrôles de réponses et PDF standard ou agrandi.",
            "Évaluation interne : V35 fournit cinq familles finies, sans modèle. Les 341 pages standard et agrandies ont été contrôlées par inspection ou comparaison intégrale avec des pages déjà inspectées. Les sujets restent plus courts que l’annale officielle comparée ; leur difficulté et leur durée ne sont pas étalonnées.",
            "Limites : aucun partenariat régional, avis d’enseignant français ni résultat auprès d’élèves n’est revendiqué. Le corpus local audité contient 79 sources indexées, dont 13 réservées à l’évaluation. Les droits de chaque document restent à vérifier.",
        ],
    )
    doc.add_heading("Présentation et aides antérieures", level=1)
    doc.add_paragraph(
        "J’ai soumis une première version britannique à une enseignante d’informatique. Ses remarques sur la difficulté des questions, le SQL, les réponses calculées, les bus informatiques et la précision des barèmes ont été transformées en contrôles et en exigences de génération."
    )
    add_label_paragraph(
        doc,
        "Financement ou demande d’aide antérieure",
        "À confirmer avant envoi.",
    )
    doc.add_heading("Adéquation au Prix Occitanie", level=1)
    doc.add_paragraph(
        "Le prix pourrait soutenir les essais sur plusieurs Mac, l’accessibilité des PDF et l’intégration de données ouvertes d’Occitanie dans des exercices liés au programme de NSI. Toute expérimentation en établissement resterait soumise à l’accord d’un partenaire. Aucune disponibilité d’enseignant ni adoption régionale n’est acquise."
    )
    doc.add_paragraph(
        "Le projet est techniquement faisable dans les douze mois parce que le moteur, le parcours macOS et la production de PDF existent déjà. Le travail restant est précisément celui que le jury peut rendre possible : confronter l’outil au territoire, mesurer son utilité et documenter aussi bien les résultats positifs que les limites."
    )

    doc.add_heading("Travaux prévus pendant les douze mois suivant le prix", level=1)
    for period, activity, evidence in (
        (
            "Mois 1 à 2",
            "Extension du catalogue et comparaison interne aux annales",
            "résultats conservés et protocole de relecture documenté",
        ),
        (
            "Mois 3 à 4",
            "Nouveaux sujets contrôlés et corrigés",
            "réponses vérifiées, défauts et révisions tracés",
        ),
        (
            "Mois 5 à 6",
            "Essais d’utilisation et étude d’un éventuel pilote",
            "parcours réels testés ; toute étude en lycée exige un accord préalable",
        ),
        (
            "Mois 7 à 9",
            "Données ouvertes d’Occitanie et audit d’accessibilité",
            "provenance, sujets relus, PDF standard et agrandi",
        ),
        (
            "Mois 10 à 12",
            "Décision sur une deuxième spécialité ou l’inférence mutualisée",
            "conclusion fondée sur les besoins établis et le matériel testé",
        ),
    ):
        add_label_paragraph(
            doc, period + ".", activity + " ; preuve : " + evidence + "."
        )
    doc.add_heading("Utilisation du prix de 1 000 euros", level=1)
    doc.add_paragraph(
        "Répartition à confirmer par le candidat avant envoi. Les usages envisagés sont les essais matériels, l’accessibilité et les données régionales. Aucun engagement de dépense ni financement d’avis d’enseignants indisponibles n’est présumé."
    )
    doc.add_heading("Faisabilité et résultats attendus", level=1)
    doc.add_paragraph(
        "Le socle logiciel fonctionne sur macOS. La justesse pédagogique reste une limite explicite. Le projet conserve les échecs et distingue contrôles techniques, relecture interne et données recueillies auprès d’élèves. Le résultat attendu est un catalogue plus varié et documenté, avec des réponses vérifiables et des PDF lisibles. Aucune recommandation d’enseignant ni amélioration mesurée des apprentissages n’est promise."
    )
    path = OUTPUT / "Prix-Occitanie-2026-Candidature-James-Durup.docx"
    OUTPUT.mkdir(parents=True, exist_ok=True)
    finalize_fonts(doc)
    doc.save(path)
    return path


def build_technical_dossier() -> Path:
    doc = Document()
    configure_document(doc, editorial=True)
    title_block(
        doc,
        "Keldin for French NSI",
        "Technical and user report Prix Occitanie 2026",
    )
    doc.add_paragraph(
        "Keldin produces non-official French NSI practice drafts. All three exercises use app-owned contracts for questions, answers and marking clauses. V35 does not call an AI model. It offers five finite whole-paper families, not unlimited new papers. They have 34 or 36 questions and 7 to 9 subject pages in standard print. Internal review compares syllabus coverage, reasoning, answers and every printed page with official material. Difficulty and duration remain uncalibrated; no teacher approval or learner trial is claimed."
    )
    doc.add_heading("The teacher’s task", level=1)
    doc.add_paragraph(
        "Teachers need additional practice without introducing errors in questions or marking. Keldin prepares a draft from its limited catalogue. The teacher decides whether to share it; the app neither marks student work nor collects student responses."
    )
    screenshot = (
        ROOT
        / "macOS/Keldin/Assets.xcassets/TutorialFrenchNSI.imageset/TutorialFrenchNSI.png"
    )
    picture = doc.add_picture(str(screenshot), width=Cm(16.2))
    picture.height = Cm(12.3)
    crop = OxmlElement("a:srcRect")
    crop.set("l", "30000")
    crop.set("r", "13000")
    crop.set("b", "20000")
    picture._inline.graphic.graphicData.pic.blipFill.insert(1, crop)
    picture._inline.docPr.set(
        "descr", "French NSI workspace in an earlier Keldin macOS prototype before V32"
    )
    caption = doc.add_paragraph(
        "Earlier prototype interface before V32. The model controls shown are not part of the current French workflow."
    )
    caption.paragraph_format.space_after = Pt(12)
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(85, 85, 85)
    doc.add_page_break()
    doc.add_heading("From request to reviewed artifact", level=1)
    for label, detail in (
        (
            "01  Prepare references.",
            "The teacher checks the local reference notices. Downloads require consent; reference access does not grant redistribution rights.",
        ),
        (
            "02  Choose a print profile.",
            "The user selects standard or enlarged print. V35 needs neither an AI model nor an Ollama installation.",
        ),
        (
            "03  Create a draft.",
            "The app selects an unused whole-paper family, constructs three exercises and checks answers and exact credits. Saved family identifiers prevent repetition across output folders.",
        ),
        (
            "04  Inspect the evidence.",
            "The teacher reads both PDFs, sources, validation results, and limitations.",
        ),
        (
            "05  Record a decision.",
            "The optional teacher-review record applies only to the exact artifact hashes. Internal engineering review is recorded separately and never creates a teacher approval.",
        ),
    ):
        add_label_paragraph(doc, label, detail)

    doc.add_heading("Architecture", level=1)
    doc.add_paragraph(
        "A SwiftUI interface exchanges JSON messages with a Python engine. The assessment registry selects an explicit French policy. This boundary prevents UK assessment objectives, integer-only marks, and English instructions from silently applying to the baccalauréat."
    )
    for label, detail in (
        (
            "Catalogue.",
            "Explicit education context, curriculum and rule versions, language, and checked compatibility.",
        ),
        (
            "References.",
            "French-only eligibility filtering; source URL, retrieval date, hash, page, rights and holdout status retained.",
        ),
        (
            "Generation.",
            "Three independent exercises with versioned facts, French prompts, proposed answers and marking clauses. Repeated families are refused.",
        ),
        (
            "Verification.",
            "Bounded SQL and algorithm checks, exact credit arithmetic and originality screening; unsupported claims remain unresolved.",
        ),
        (
            "Publication and review.",
            "Atomic PDF bundle plus manifest; optional external review records are bound to file hashes. Unavailable reviews remain not run.",
        ),
    ):
        add_label_paragraph(doc, label, detail)

    doc.add_heading("Implemented French assessment rules", level=1)
    add_bullets(
        doc,
        [
            "Terminale written component, session 2027: 3 hours 30 minutes and three independent exercises.",
            "Eighteen technical points plus a separate two-point French-language component. Detailed credit is indicative, not an official marking grid.",
            "Première knowledge may be a prerequisite but cannot replace the essential Terminale focus.",
            "Each question records one to three official capability codes, a cognitive operation, a 1–4 difficulty label, and an estimated time.",
            "Each exercise covers its declared topics, totals 70 minutes, uses at least three cognitive operations, and includes a difficulty-4 question.",
        ],
    )

    doc.add_heading("Question quality", level=1)
    doc.add_paragraph(
        "The five families combine graph or tree reasoning, SQL and network or process cases. The new family connects object graphs and breadth-first search, eligibility joins and backtracking allocation, and a specified blocking communication channel. Candidates construct helpers, trace state changes and explain counterexamples. Package replay checks versioned prompts, answers and credits. Five families remain limited variety, not full curriculum coverage."
    )
    doc.add_paragraph(
        "The schema rejects recall-only sequences and requires progression towards analysis, design, debugging or justification. These declared labels do not measure learner difficulty. Calibration still needs comparison with held-out papers and observed completion times."
    )

    doc.add_heading("PDF fidelity", level=1)
    doc.add_paragraph(
        "The A4 profile was measured against a 2026 Métropole paper: 595.32 × 841.92 points, with 12-point body text and monospaced code. V35 standard subjects have 7 to 9 pages, against the official comparator’s 16 pages and 36 questions. The new family has 36 questions, a nine-page subject and nineteen-page correction. All 341 standard/enlarged pages were inspected or matched in full-page rasters to previously inspected pages. Source-content, graph-arrow and nonempty-page checks supplement visual review. Sparse tails remain. Page counts do not establish equivalent reasoning depth, difficulty or working time."
    )

    doc.add_heading("Safety, privacy, and rights", level=1)
    for label, detail in (
        (
            "Data exposure.",
            "Public V35 French generation makes no AI-provider call and collects no student responses. Historical model experiments are separate; there is no silent cloud fallback.",
        ),
        (
            "Generated code.",
            "SQL runs in an isolated, bounded setting; supported code has restricted interpretation. Arbitrary generated Python is never executed.",
        ),
        (
            "Wrong source.",
            "Scope filtering precedes ranking; holdouts are separate and an empty eligible set fails explicitly. Third-party rights still need document-level review.",
        ),
        (
            "Past-paper copying.",
            "Text, normalised code, structure and prior generations are compared. Similarity is risk evidence, not a legal guarantee.",
        ),
        (
            "Document rights.",
            "Local references retain URL, rights status and hash. Third-party illustrations require separate clearance.",
        ),
    ):
        add_label_paragraph(doc, label, detail)

    doc.add_page_break()
    doc.add_heading("Qualification evidence", level=1)
    doc.add_paragraph(
        "The historical model benchmark tests ten papers per model at fixed seeds and retains failures and identities. It is separate from the public V35 catalogue. Repeating a family or rendering it in enlarged print does not add an independent paper to a qualification sample."
    )
    add_label_paragraph(
        doc,
        "Initial campaign · 0/10 accepted.",
        "Gemma 4 12B passed 0/10 in the first campaign. Those failures remain historical model evidence. V35 has no model to qualify. Source, packaged CLI and native app checks cover five families and repeat refusal. Internal review found no blocking defect in accepted artifacts, without establishing classroom suitability or calibrated difficulty.",
    )
    add_label_paragraph(
        doc,
        "Internal review and external limits.",
        "Teacher review is unavailable. Internal review checks the programme, task progression, independent answers, marking and page renders. Teacher-review and empirical-calibration records remain not run; no approval is invented.",
    )
    add_label_paragraph(
        doc,
        "Learners · not piloted.",
        "Learner timing, accessibility and teacher workload remain unmeasured. A future pilot needs a partner and consent; no school is arranged.",
    )

    doc.add_heading("Occitanie deployment", level=1)
    doc.add_paragraph(
        "Users generate on a Mac and share ordinary PDFs. Occitanie’s loRdi devices can display PDFs; these Windows laptops are not claimed to run the app or a large model. A regional trial is a possible future extension, with no school agreement or deployment claimed."
    )
    doc.add_paragraph(
        "A planned regional module will use rights-cleared open datasets on schools, energy or transport. Official values will remain distinct from synthetic exercise data. Each context must support an NSI capability such as SQL or graph algorithms; place names alone add no educational value."
    )

    doc.add_heading("Risks and decisions", level=1)
    for label, detail in (
        (
            "Wrong answer.",
            "Block publication, repair the source and repeat answer and PDF checks.",
        ),
        (
            "Memory pressure.",
            "Benchmark actual Macs; promise no universal compatibility.",
        ),
        ("2027 visual rules.", "Keep the historical-reference layout provisional."),
        (
            "Regional relevance.",
            "Test useful data and a real pilot, not renamed places.",
        ),
        ("Unfavourable pilot.", "Report limits, narrow scope, or stop."),
    ):
        add_label_paragraph(doc, label, detail)

    doc.add_heading("Primary sources", level=1)
    sources = [
        (
            "Prix Occitanie 2026 rules",
            "https://association.centralesupelec-alumni.com/medias/editor/PRIX_OCCITANIE_2026/REGLEMENT_PRIX_OCCITANIE_2026.pdf",
        ),
        (
            "Official NSI 2027 examination definition",
            "https://www.education.gouv.fr/bo/2026/Special4/MENE2622643N",
        ),
        (
            "NSI curriculum and resources",
            "https://eduscol.education.gouv.fr/5823/programmes-et-ressources-en-numerique-et-sciences-informatiques-voie-g",
        ),
        (
            "Official past-paper archive",
            "https://eduscol.education.gouv.fr/5199/annales-des-epreuves-du-baccalaureat-des-voies-generale-et-technologique",
        ),
        (
            "loRdi",
            "https://www.laregion.fr/aide-a-l-acquisition-d-un-ordinateur-portable-lordi",
        ),
        (
            "Occitanie school strategy",
            "https://www.laregion.fr/Quelle-demarche-pour-le-lycee-de-demain-39604",
        ),
        (
            "Occitanie schools open data",
            "https://data.laregion.fr/explore/dataset/lycees-occitanie/",
        ),
        (
            "Occitanie energy-use history",
            "https://data.laregion.fr/explore/dataset/historique-de-la-consommation-denergie-en-occitanie/",
        ),
        (
            "CNIL guidance for teachers",
            "https://www.cnil.fr/fr/enseignant-usage-systeme-ia",
        ),
    ]
    for index, (label, url) in enumerate(sources, 1):
        add_source(doc, index, label, url)

    path = OUTPUT / "Keldin-NSI-Technical-and-User-Report.docx"
    OUTPUT.mkdir(parents=True, exist_ok=True)
    finalize_fonts(doc)
    doc.save(path)
    return path


def build_mathematical_analysis() -> Path:
    doc = Document()
    configure_document(doc, compact=True, editorial=True)
    title_block(
        doc,
        "Mathematical analysis of qualification",
        "Keldin Occitanie NSI review 28 October 2026",
    )
    doc.add_paragraph(
        "Public V35 supplies five finite paper families without an AI-model call. Repeated families or enlarged copies are not independent new samples. Internal review cannot prove psychometric equivalence to the baccalauréat. Learner completion times and success rates have not been measured; the statistical examples are illustrative."
    )
    doc.add_paragraph(
        "The first campaign on 3 October 2026 accepted 0/10 complete Gemma 4 12B papers. The hypothetical zero-defect calculations below do not apply to that campaign or justify a model recommendation."
    )

    doc.add_heading("Exact paper constraints", level=1)
    doc.add_paragraph(
        "For three exercises e = 1, 2, 3, the plan assigns mₑ = 70 estimated minutes each, totalling 210 minutes or 3 hours 30 minutes. The public catalogue allocates 6 points per exercise: Σpₑ = 18. Separate language credit adds exactly 2 points, making 20. Credit uses exact decimal arithmetic. The allocation is indicative, not an official marking grid."
    )
    doc.add_paragraph(
        "For every question q, its marking criteria must satisfy Σcᵢ = p_q; estimated question times in each exercise must satisfy Σt_q = 70 minutes. A mismatch—even 0.5 point—blocks the document."
    )

    doc.add_heading("Cognitive demand", level=1)
    doc.add_paragraph(
        "Each question is labelled as recall, application, analysis, design, debugging, or justification, and receives an integer difficulty label from 1 to 4. A valid exercise has at most one recall-only question, at least three different operations, at least two higher-order operations, and at least one difficulty-4 question. These constraints are structural guardrails; they do not establish how difficult learners will actually find the work."
    )

    doc.add_heading("Defect rates and sample size", level=1)
    doc.add_paragraph(
        "Let p be the probability that an exercise has a blocking defect detectable during review. Given n independent exercises and zero observed defects, a simple one-sided 95% upper bound is p₉₅ = 1 − 0.05^(1/n). For n = 30 exercises from one model, p₉₅ ≈ 9.5%. Even a flawless 30-exercise run would not establish a true defect rate below 1%; roughly 299 defect-free exercises would be needed for that bound."
    )
    doc.add_paragraph(
        "Shared prompts or model failures weaken independence; defect rates must therefore be reported by category and version."
    )
    sample_table = add_table(
        doc,
        ["Zero-defect sample", "95% upper bound", "Interpretation"],
        [
            ["30 exercises", "9.5%", "Initial model screen only."],
            ["90 exercises", "3.3%", "Pooling cannot prove each model."],
            ["299 exercises", "1.0%", "Approximate 1% scale."],
        ],
        [4.0, 4.0, 8.8],
    )
    # This four-row comparison should remain readable as one unit.
    for row in sample_table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    doc.add_heading("Internal review and unavailable evidence", level=1)
    doc.add_paragraph(
        "Internal review checks curriculum alignment, correctness, French, sustained reasoning, marking, layout and originality against official material. Independent executable cases supplement answer review where possible. All five families receive source and page checks; a high average cannot hide a blocking error. This is engineering evidence, not teacher approval or observed learner performance."
    )
    doc.add_paragraph(
        "No teacher panel is available. Teacher review and empirical calibration remain not run. Reviewer agreement statistics such as weighted kappa require real independent ratings; none are calculated or implied here. Repeated checks by one reviewer do not create independent opinions."
    )

    doc.add_heading("Timing calibration", level=1)
    doc.add_paragraph(
        "For learner i and paper j, relative duration error is eᵢⱼ = (Tᵢⱼ − 210) / 210. The median T better describes a typical completion time than a mean distorted by abandoned attempts. Results will include the median, interquartile range, number of completed papers, and share exceeding 210 minutes. Contexts and approved accommodations will be reported separately; a small pilot cannot define a national norm."
    )

    doc.add_heading("Teacher time", level=1)
    doc.add_paragraph(
        "Time saved per paper is Δt = t_manual − (t_supervised generation + t_review + t_repair). A future pilot would need to measure each term rather than ask for a general impression. Across N papers, annual time saved is N × Δt. Illustration only: if a manual paper takes 120 minutes and assisted generation, review, and repair take 35, twelve papers save 17 hours. This is not an observed result."
    )

    doc.add_heading("Layout fidelity", level=1)
    doc.add_paragraph(
        "Visual comparison uses measurements, not an opaque ‘similarity percentage’. For feature k—such as a margin, font size, or heading position—the normalised distance is d_k = |x_k − r_k| / s_k, where r_k is the reference measurement and s_k is a justified tolerance. Aggregate layout distance D = Σw_k d_k flags regressions; published weights sum to 1. This score says nothing about question quality."
    )
    doc.add_paragraph(
        "The current profile uses A4 at 595.32 × 841.92 points, 12-point body text, a left margin near 70.6 points, 14-point exercise titles, and cover positions measured from a 2026 Métropole paper. Independent branding and the non-official notice are deliberate differences."
    )

    doc.add_heading("Originality", level=1)
    doc.add_paragraph(
        "Screening compares text n-grams, normalised code, question structure, and previous generations. Thresholds must be calibrated on labelled copies, superficial renamings, and genuinely different tasks; precision and recall are reported separately. A threshold is acceptable only if it rejects copied and renamed holdout cases without discarding too many original ones. This reduces risk but is not a copyright guarantee."
    )

    doc.add_heading("Energy and resource use", level=1)
    doc.add_paragraph(
        "Model download size is neither memory use nor energy consumption. During benchmarking, energy per accepted paper will be E = P_mean × t / 3600, with P in watts and t in seconds. The report will include energy spent on rejected attempts and the number of times a PDF is reused. No carbon-reduction claim is justified without measured power, an explicit emissions factor, and a relevant comparator."
    )

    doc.add_heading("Budget and decision rule", level=1)
    for index, (label, detail) in enumerate(
        (
            (
                "Budget subject to confirmation.",
                "The applicant must confirm any allocation. No teacher availability or paid review is assumed.",
            ),
            (
                "Possible school trial.",
                "Requires an agreed partner and actual access conditions.",
            ),
            (
                "Hardware and accessibility.",
                "Tested configurations, memory, speed, and standard/enlarged PDFs.",
            ),
            (
                "Data and contingency.",
                "Regional resources or a documented pilot need.",
            ),
        )
    ):
        add_label_paragraph(doc, label, detail)
        doc.paragraphs[-1].paragraph_format.keep_with_next = index < 3
    doc.add_paragraph(
        "An engineering release requires no known blocking defect, reproducible source and package checks, readable reviewed PDFs and working native actions on the tested Mac. Educational effectiveness, plausible completion time and teacher time savings remain unmeasured. A future school deployment would need its own evidence and approval; passing engineering checks cannot supply them."
    )

    doc.add_heading("References", level=1)
    add_source(
        doc,
        1,
        "NSI 2027 rules",
        "https://www.education.gouv.fr/bo/2026/Special4/MENE2622643N",
    )
    add_source(
        doc,
        2,
        "Official past papers",
        "https://eduscol.education.gouv.fr/5199/annales-des-epreuves-du-baccalaureat-des-voies-generale-et-technologique",
    )
    add_source(
        doc,
        3,
        "Official French-language rubric",
        "https://www.education.gouv.fr/sites/default/files/document/annexe-attendus-et-observables-redactionnels-520693.pdf",
    )
    add_source(
        doc,
        4,
        "Prix Occitanie 2026 rules",
        "https://association.centralesupelec-alumni.com/medias/editor/PRIX_OCCITANIE_2026/REGLEMENT_PRIX_OCCITANIE_2026.pdf",
    )

    path = OUTPUT / "Keldin-NSI-Mathematical-Analysis.docx"
    OUTPUT.mkdir(parents=True, exist_ok=True)
    finalize_fonts(doc)
    doc.save(path)
    return path


def main(argv: list[str] | None = None) -> None:
    global OUTPUT
    parser = argparse.ArgumentParser(
        description="Build the maintained Prix Occitanie 2026 application evidence documents."
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=OUTPUT,
        help="Output directory. Use a fresh directory to preserve completed personal forms.",
    )
    OUTPUT = parser.parse_args(argv).output.expanduser().resolve()
    for path in (
        build_application(),
        build_technical_dossier(),
        build_mathematical_analysis(),
    ):
        print(path)


if __name__ == "__main__":
    main()
