import pytest
from docx import Document

from tools import build_occitanie_submission as submission
from tools.build_occitanie_submission import (
    build_application,
    build_mathematical_analysis,
    build_technical_dossier,
    configure_document,
    finalize_fonts,
    title_block,
)


# Tests must not overwrite the applicant's editable submission copies.
@pytest.fixture(autouse=True)
def isolate_generated_submission_files(tmp_path, monkeypatch):
    monkeypatch.setattr(submission, "OUTPUT", tmp_path)


def test_application_title_uses_a_real_title_paragraph() -> None:
    document = Document()
    configure_document(document, compact=True)

    title_block(document, "Candidature au Prix Occitanie 2026", "Keldin")

    assert document.paragraphs[0].style.name == "Title"
    assert document.paragraphs[0].text == "Candidature au Prix Occitanie 2026"
    assert document.paragraphs[1].text == "Keldin"
    assert not document.tables


def test_explicit_fonts_preserve_the_document_hierarchy() -> None:
    document = Document()
    configure_document(document, compact=True)
    title_block(document, "Candidature au Prix Occitanie 2026", "Keldin")
    heading = document.add_heading("Présentation du candidat", level=1)
    heading.add_run(" ")

    finalize_fonts(document)

    assert document.paragraphs[0].runs[0].font.size.pt == 22
    assert document.paragraphs[0].runs[0].font.name == "Arial"
    assert heading.runs[0].font.size.pt == 13.5


def test_supporting_reports_are_in_english_while_the_application_stays_french() -> None:
    application = Document(build_application())
    technical = Document(build_technical_dossier())
    mathematical = Document(build_mathematical_analysis())

    assert application.paragraphs[0].text == "Candidature au Prix Occitanie 2026"
    assert technical.paragraphs[0].text == "Keldin for French NSI"
    assert mathematical.paragraphs[0].text == "Mathematical analysis of qualification"
    assert "0/10 accepted" in " ".join(
        paragraph.text for paragraph in technical.paragraphs
    )
    assert "five finite whole-paper families" in " ".join(
        paragraph.text for paragraph in technical.paragraphs
    )


def test_technical_report_distinguishes_current_catalogue_from_historical_model_trials() -> (
    None
):
    technical = Document(build_technical_dossier())
    text = " ".join(paragraph.text for paragraph in technical.paragraphs).lower()

    assert "all three exercises use app-owned contracts" in text
    assert "v35 does not call an ai model" in text
    assert "five finite whole-paper families" in text
    assert "difficulty and duration remain uncalibrated" in text
    assert "34 or 36 questions" in text
    assert "7 to 9 subject pages" in text
    assert "nineteen-page correction" in text
    assert "historical model" in text
    assert "choose a model" not in text
    assert "v25 single-paper diagnostic" not in text
    assert "18 numbered questions across seven subject pages" not in text
    assert "eleven-page proposed correction" not in text
    assert "other two exercises remain model-authored" not in text
    assert "has not passed a live paper" not in text


def test_current_submission_brand_and_french_scope_are_explicit():
    for builder in (
        build_application,
        build_technical_dossier,
        build_mathematical_analysis,
    ):
        text = " ".join(p.text for p in Document(builder()).paragraphs)
        assert "Keldin" in text
        assert "Paper Creator" not in text
        assert "V35" in text
    application = " ".join(p.text for p in Document(build_application()).paragraphs)
    assert "L’enseignant choisit un modèle local" not in application
    assert "cinq familles" in application


def test_reports_state_v35_scope_and_do_not_require_unavailable_teacher_reviews():
    application = " ".join(p.text for p in Document(build_application()).paragraphs)
    technical = " ".join(p.text for p in Document(build_technical_dossier()).paragraphs)
    mathematical = " ".join(
        p.text for p in Document(build_mathematical_analysis()).paragraphs
    )
    assert "V35" in application and "cinq familles" in application
    assert "five finite whole-paper families" in technical
    assert "34 or 36 questions" in technical
    assert "7 to 9 subject pages" in technical
    assert "internal review" in technical.lower()
    assert "not been measured" in mathematical
    for text in (technical, mathematical):
        assert "Two NSI teachers must" not in text
        assert "Two teachers will score" not in text
        assert "two written teacher recommendations" not in text
        assert "V32 supplies" not in text


def test_mathematical_report_uses_current_equal_exercise_credits():
    text = " ".join(p.text for p in Document(build_mathematical_analysis()).paragraphs)
    assert "6 points per exercise" in text
    assert "5.5, 6, and 6.5" not in text


def test_submission_can_be_built_without_replacing_an_existing_personal_form(tmp_path):
    original = tmp_path / "Prix-Occitanie-2026-Candidature-James-Durup.docx"
    original.write_bytes(b"Applicant-owned completed form")
    destination = tmp_path / "refreshed"
    submission.main(["--output", str(destination)])
    assert original.read_bytes() == b"Applicant-owned completed form"
    assert len(list(destination.glob("*.docx"))) == 3


def test_technical_report_shows_and_describes_the_real_mac_workflow() -> None:
    document = Document(build_technical_dossier())

    assert len(document.inline_shapes) == 1
    picture = document.inline_shapes[0]._inline.docPr
    assert "French NSI" in picture.get("descr", "")
    crop = document.element.body.xpath(".//a:srcRect")
    assert len(crop) == 1
    assert int(crop[0].get("l")) >= 20000
    assert int(crop[0].get("b")) >= 20000
    assert document.inline_shapes[0].height.cm >= 12
    assert document.inline_shapes[0].height.cm <= 13
    assert any(
        "prototype interface" in paragraph.text.lower()
        for paragraph in document.paragraphs
    )


def test_supporting_reports_have_a_distinct_editorial_hierarchy() -> None:
    application = Document(build_application())
    technical = Document(build_technical_dossier())
    mathematical = Document(build_mathematical_analysis())

    assert application.paragraphs[0].runs[0].font.size.pt == 22
    for document in (technical, mathematical):
        assert document.paragraphs[0].runs[0].font.size.pt >= 28
    assert technical.styles["Normal"].font.size.pt >= 11
    assert mathematical.styles["Normal"].font.size.pt >= 10.5


def test_supporting_sources_are_clickable_without_long_printed_urls() -> None:
    for path, expected_links in (
        (build_technical_dossier(), 9),
        (build_mathematical_analysis(), 4),
    ):
        document = Document(path)
        links = [
            relation.target_ref
            for relation in document.part.rels.values()
            if relation.reltype.endswith("/hyperlink")
        ]
        assert len(links) == expected_links
        assert all(link.startswith("https://") for link in links)
        assert not any(
            "https://" in paragraph.text for paragraph in document.paragraphs
        )


def test_editorial_reports_do_not_break_large_grids_across_pages() -> None:
    technical = Document(build_technical_dossier())
    mathematical = Document(build_mathematical_analysis())

    assert not technical.tables
    assert len(mathematical.tables) == 1
    assert not mathematical.element.body.xpath(".//w:br[@w:type='page']")


def test_formal_application_uses_readable_milestones_instead_of_split_grids() -> None:
    application = Document(build_application())
    text = " ".join(paragraph.text for paragraph in application.paragraphs)

    assert not application.tables
    assert "Mois 10 à 12" in text
    assert "Répartition à confirmer" in text
    assert "À compléter" in text
    assert "difficulté et leur durée ne sont pas étalonnées" in text
